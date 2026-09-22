from collections import defaultdict
from datetime import UTC, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .service import ROLE_SECTIONS, display_role, minutes_between


NAVY = '17263D'
TEAL = '087F8C'
PALE = 'EAF6F4'
ORANGE = 'FFF1DC'
GRAY = 'F3F6FA'
WHITE = 'FFFFFF'
TABLE_BORDER = Border(
    left=Side(style='thin', color='BCCAD5'),
    right=Side(style='thin', color='BCCAD5'),
    top=Side(style='thin', color='BCCAD5'),
    bottom=Side(style='thin', color='BCCAD5'),
)


def safe_text(value):
    text = str(value or '')
    return "'" + text if text.startswith(('=', '+', '-', '@')) else text


def as_utc(value):
    if value is None:
        return None
    return value.replace(tzinfo=UTC) if value.tzinfo is None else value.astimezone(UTC)


def minutes_for(session, now):
    return session.duration_minutes if session.duration_minutes is not None else minutes_between(session.start_time, now)


def section_bar(sheet, row, title, columns):
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=columns)
    cell = sheet.cell(row, 1, title)
    cell.fill = PatternFill('solid', fgColor=TEAL)
    cell.font = Font(color=WHITE, bold=True, size=11)
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.border = TABLE_BORDER
    sheet.row_dimensions[row].height = 25


def heading(sheet, row, labels):
    for column, label in enumerate(labels, 1):
        cell = sheet.cell(row, column, label)
        cell.fill = PatternFill('solid', fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = TABLE_BORDER
    sheet.row_dimensions[row].height = 23


def data_row(sheet, row, values, alternate=False):
    for column, value in enumerate(values, 1):
        cell = sheet.cell(row, column, safe_text(value) if isinstance(value, str) else value)
        cell.fill = PatternFill('solid', fgColor=GRAY if alternate else WHITE)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = TABLE_BORDER
    sheet.row_dimensions[row].height = 21


def total_row(sheet, row, label, count, minutes, columns=7):
    data_row(sheet, row, [label, f'{count} breaks', '', minutes] + [''] * (columns - 4))
    for cell in sheet[row][:columns]:
        cell.fill = PatternFill('solid', fgColor=PALE)
        cell.font = Font(color=NAVY, bold=True)
    sheet.row_dimensions[row].height = 23


def export_break_workbook(users, sessions, date_from, date_to):
    workbook = Workbook()
    summary = workbook.active
    summary.title = 'Summary'
    details = workbook.create_sheet('Break history')
    now = datetime.now(UTC)
    by_user = defaultdict(list)
    for session in sessions:
        by_user[session.user_id].append(session)

    for sheet, columns in ((summary, 5), (details, 7)):
        sheet.sheet_view.showGridLines = False
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
        title = sheet.cell(1, 1, 'EMPLOYEE BREAK REPORT')
        title.fill = PatternFill('solid', fgColor=NAVY)
        title.font = Font(color=WHITE, bold=True, size=17)
        sheet.row_dimensions[1].height = 36
        sheet.cell(2, 1, f'Period: {date_from:%d %b %Y} - {date_to:%d %b %Y}  |  Times in UTC')
        sheet.cell(2, 1).font = Font(color='53657A', italic=True)
        sheet.row_dimensions[2].height = 25

    summary.cell(4, 1, 'Employees')
    summary.cell(4, 2, len(users))
    summary.cell(4, 3, 'Breaks')
    summary.cell(4, 4, len(sessions))
    summary.cell(5, 1, 'Total minutes')
    summary.cell(5, 2, sum(minutes_for(item, now) for item in sessions))
    for row in (4, 5):
        for cell in summary[row][:4]:
            cell.fill = PatternFill('solid', fgColor=PALE)
            cell.font = Font(color=NAVY, bold=True)
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = TABLE_BORDER

    summary_row = 7
    detail_row = 4
    for role_key, role_title, _ in ROLE_SECTIONS:
        role_users = [user for user in users if display_role(user) == role_key]
        if not role_users:
            continue
        section_bar(summary, summary_row, role_title, 5)
        section_bar(details, detail_row, role_title, 7)
        summary_row += 1
        detail_row += 1
        heading(summary, summary_row, ['Employee', 'Role', 'Breaks', 'Total minutes', 'Status'])
        summary_row += 1

        for user in role_users:
            entries = by_user[user.id]
            total = sum(minutes_for(item, now) for item in entries)
            active = any(item.end_time is None for item in entries)
            data_row(summary, summary_row, [user.full_name, role_title, len(entries), total, 'On break' if active else 'Active'], summary_row % 2 == 0)
            if active:
                summary.cell(summary_row, 5).fill = PatternFill('solid', fgColor=ORANGE)
            summary_row += 1

            details.merge_cells(start_row=detail_row, start_column=1, end_row=detail_row, end_column=7)
            employee = details.cell(detail_row, 1, safe_text(user.full_name))
            employee.fill = PatternFill('solid', fgColor=PALE)
            employee.font = Font(color=NAVY, bold=True)
            employee.alignment = Alignment(horizontal='center', vertical='center')
            employee.border = TABLE_BORDER
            details.row_dimensions[detail_row].height = 23
            detail_row += 1
            if not entries:
                data_row(details, detail_row, ['No breaks in selected range', '', '', '', '', '', ''])
                detail_row += 1
            by_date = defaultdict(list)
            for session in entries:
                by_date[as_utc(session.start_time).date()].append(session)
            for break_date, day_entries in sorted(by_date.items()):
                section_bar(details, detail_row, break_date.strftime('%A, %d %B %Y'), 7)
                detail_row += 1
                heading(details, detail_row, ['Date', 'Start (UTC)', 'End (UTC)', 'Duration (mins)', 'Reason', 'Status', 'Employee'])
                detail_row += 1
                for session in day_entries:
                    start = as_utc(session.start_time)
                    end = as_utc(session.end_time)
                    data_row(details, detail_row, [start.date(), start.strftime('%H:%M:%S'), end.strftime('%H:%M:%S') if end else '', minutes_for(session, now), session.reason, 'Completed' if end else 'In progress', user.full_name], detail_row % 2 == 0)
                    details.cell(detail_row, 1).number_format = 'dd mmm yyyy'
                    if not end:
                        details.cell(detail_row, 6).fill = PatternFill('solid', fgColor=ORANGE)
                    detail_row += 1
                total_row(details, detail_row, 'Daily total', len(day_entries), sum(minutes_for(item, now) for item in day_entries))
                detail_row += 2
            total_row(details, detail_row, 'Employee total', len(entries), total)
            detail_row += 1
            detail_row += 1
        summary_row += 1
        detail_row += 1

    if not users:
        data_row(summary, 7, ['No employees found', '', '', '', ''])
        data_row(details, 4, ['No break records found', '', '', '', '', '', ''])

    for sheet, widths in ((summary, [30, 20, 13, 19, 16]), (details, [18, 18, 18, 21, 28, 18, 30])):
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.freeze_panes = 'A8' if sheet == summary else 'A4'
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
